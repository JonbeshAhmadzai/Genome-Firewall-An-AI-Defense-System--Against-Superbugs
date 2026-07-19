#!/usr/bin/env python3
"""Build the short Genome Firewall team-video deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path("docs/team_video_presentation.pptx")
NAVY = RGBColor(11, 19, 43)
PANEL = RGBColor(20, 33, 61)
CYAN = RGBColor(54, 197, 240)
GOLD = RGBColor(255, 204, 102)
WHITE = RGBColor(244, 247, 251)
MUTED = RGBColor(172, 190, 211)


def text_box(slide, text, x, y, w, h, size=20, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def panel(slide, heading, copy, y, accent=GOLD):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(y), Inches(12.0), Inches(0.82))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.fill.background()
    text_box(slide, heading, 0.95, y + 0.08, 2.45, 0.58, size=17, color=accent, bold=True)
    text_box(slide, copy, 3.35, y + 0.08, 9.0, 0.58, size=16, color=WHITE)


def base_slide(prs, title, timing, speaker):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = CYAN; bar.line.fill.background()
    text_box(slide, title, 0.7, 0.42, 11.9, 0.58, size=27, color=WHITE, bold=True)
    text_box(slide, f"{timing}  ·  {speaker.upper()}", 0.72, 1.05, 8.0, 0.28, size=11, color=CYAN, bold=True)
    text_box(slide, "GENOME FIREWALL  /  RESEARCH PROTOTYPE", 0.7, 7.12, 8.0, 0.2, size=9, color=MUTED, bold=True)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = base_slide(prs, "Genome Firewall", "0:00–0:06", "All speakers")
    text_box(slide, "Transparent genomic AI\nfor antibiotic-resistance evidence", 0.75, 1.85, 7.0, 1.6, size=31, color=WHITE, bold=True)
    text_box(slide, "Evidence first  ·  uncertainty visible  ·  human oversight", 0.78, 3.72, 7.2, 0.45, size=18, color=CYAN)
    panel(slide, "SAFETY", "Confirm every result with standard laboratory testing.", 5.25, accent=GOLD)

    slide = base_slide(prs, "1 · Data pipeline", "0:06–0:18", "Speaker 1 · Data")
    panel(slide, "INPUT", "BV-BRC bacterial genomes and reconstructed FASTA files", 1.75)
    panel(slide, "PROCESS", "Cohort selection, FASTA quality control, and reproducible metadata", 2.85)
    panel(slide, "OUTPUT", "A clean, traceable cohort ready for annotation and modeling", 3.95, accent=CYAN)
    text_box(slide, "Speaker line: I organized the data so every downstream result can be reproduced.", 0.75, 5.55, 11.8, 0.48, size=17, color=MUTED, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "2 · Feature engineering", "0:18–0:30", "Speaker 2 · Features")
    panel(slide, "EVIDENCE", "AMRFinderPlus genes, mutations, and resistance determinants", 1.75)
    panel(slide, "FEATURES", "Presence flags, target signals, and engineered evidence counts", 2.85)
    panel(slide, "MATRIX", "115 reproducible features shared consistently by every model", 3.95, accent=CYAN)
    text_box(slide, "FASTA  →  AMRFinderPlus  →  feature matrix", 1.0, 5.65, 11.3, 0.42, size=20, color=GOLD, bold=True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "3 · Machine learning", "0:30–0:45", "Speaker 3 · ML")
    panel(slide, "VALIDATION", "Group-disjoint 60% train · 15% calibration · 25% test", 1.55)
    panel(slide, "MODELS", "Calibrated models plus logistic regression, random forest, extra trees, and XGBoost", 2.65)
    panel(slide, "DECISION", "Confidence threshold, target compatibility, and a transparent no-call policy", 3.75, accent=CYAN)
    text_box(slide, "Genetic groups stay separated to reduce leakage.", 0.9, 5.45, 11.6, 0.55, size=18, color=MUTED, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, "4 · UI and dashboard", "0:45–0:58", "Speaker 4 · Product")
    panel(slide, "PRODUCT", "FastAPI JSON backend, HTML dashboard, model comparison, and Streamlit fallback", 1.75)
    panel(slide, "EXPLANATION", "Optional GPT summary explains structured evidence but cannot change predictions", 2.85)
    panel(slide, "CLOSE", "A result people can inspect: QC, prediction, confidence, evidence, and safety warning", 3.95, accent=CYAN)
    text_box(slide, "Evidence first. Uncertainty visible.", 1.0, 5.55, 11.3, 0.52, size=24, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "CONFIRM EVERY RESULT WITH STANDARD LABORATORY AST", 1.0, 6.35, 11.3, 0.3, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
